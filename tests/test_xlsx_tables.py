"""Fictional, nonpublishing structured XLSX source-cell projection."""

from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import Workbook
from test_sandbox_runner import CID, IMAGE, fake_docker

from masterbrain.datacore.document_contract import (
    SHEET_TYPE,
    DocumentUnavailable,
    encode_request,
    source_for,
)
from masterbrain.datacore.document_worker import extract_document
from masterbrain.datacore.xlsx_tables import (
    decode_projection,
    encode_projection,
    project_xlsx_tables,
    require_indexed_text_digest,
)
from masterbrain.datacore.xlsx_tables_processor import DockerXlsxTableProcessor
from masterbrain.datacore.xlsx_tables_worker import process_request


def _workbook(*, formula: bool = False, rows: int = 2, hidden: bool = False) -> bytes:
    book = Workbook()
    first = book.active
    first.title = "Rent roll"
    first.append(["Project", "Monthly USD", "Note"])
    first.append(["Fictional-0", 100, None])
    first.append([None, None, None])
    for index in range(1, rows):
        first.append([f"Fictional-{index}", index + 100, "review"])
    if formula:
        first["B2"] = "=50+50"
    if hidden:
        second = book.create_sheet("Hidden source")
        second.sheet_state = "hidden"
        second["A1"] = "synthetic hidden note"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _source(raw: bytes):
    digest = hashlib.sha256(raw).hexdigest()
    return source_for(raw, "version_" + digest[:32], "rent.xlsx", digest, SHEET_TYPE)


def test_rows_retain_physical_coordinates_blanks_and_hidden_sheet_state() -> None:
    raw = _workbook(rows=3, hidden=True)
    result = project_xlsx_tables(raw, _source(raw))
    assert result.status == "extracted" and result.reason == "bounded_xlsx_tables_v1"
    assert result.scope == "bounded_textual_cell_values_only"
    assert [(sheet.index, sheet.name, sheet.visibility) for sheet in result.sheets] == [
        (1, "Rent roll", "visible"), (2, "Hidden source", "hidden"),
    ]
    assert [(row.number, row.values) for row in result.sheets[0].rows] == [
        (1, ("Project", "Monthly USD", "Note")),
        (2, ("Fictional-0", "100", "")),
        (4, ("Fictional-1", "101", "review")),
        (5, ("Fictional-2", "102", "review")),
    ]
    assert result.sheets[1].rows[0].values == ("synthetic hidden note",)


def test_thousand_row_tail_is_addressable_not_truncated() -> None:
    raw = _workbook(rows=1000)
    result = project_xlsx_tables(raw, _source(raw))
    assert result.status == "extracted"
    assert result.sheets[0].rows[-1].number == 1002
    assert result.sheets[0].rows[-1].values == ("Fictional-999", "1099", "review")


@pytest.mark.parametrize("hidden", [False, True])
def test_cells_bind_to_the_exact_existing_indexed_document_digest(hidden: bool) -> None:
    raw = _workbook(rows=3, hidden=hidden)
    source = _source(raw)
    indexed = extract_document(raw, source)
    projection = project_xlsx_tables(raw, source)
    assert indexed.status == "extracted" and projection.status == "extracted"
    digest = hashlib.sha256(indexed.text.encode()).hexdigest()
    assert require_indexed_text_digest(projection, digest) == digest
    with pytest.raises(DocumentUnavailable):
        require_indexed_text_digest(projection, "0" * 64)
    changed_row = replace(
        projection.sheets[0].rows[1],
        values=("Fictional-0", "999", ""),
    )
    changed_sheet = replace(
        projection.sheets[0],
        rows=(projection.sheets[0].rows[0], changed_row, *projection.sheets[0].rows[2:]),
    )
    forged = replace(projection, sheets=(changed_sheet, *projection.sheets[1:]))
    with pytest.raises(DocumentUnavailable):
        require_indexed_text_digest(forged, digest)
    # The indexed text does not encode physical row coordinates. The pinned
    # worker/source identity, not this text digest alone, attests positions.
    moved = replace(projection.sheets[0].rows[1], number=3)
    moved_sheet = replace(
        projection.sheets[0],
        rows=(projection.sheets[0].rows[0], moved, *projection.sheets[0].rows[2:]),
    )
    assert require_indexed_text_digest(
        replace(projection, sheets=(moved_sheet, *projection.sheets[1:])), digest,
    ) == digest


def test_formula_or_changed_bytes_return_no_partial_cells() -> None:
    raw = _workbook(formula=True)
    formula = project_xlsx_tables(raw, _source(raw))
    assert formula.status == "deferred" and formula.reason == "unsupported_features"
    assert formula.sheets == () and formula.cell_count == 0
    with pytest.raises(DocumentUnavailable):
        require_indexed_text_digest(formula, "0" * 64)
    changed = project_xlsx_tables(raw + b"changed", _source(raw))
    assert changed.status == "failed" and changed.sheets == ()
    wrong_digest = project_xlsx_tables(raw, replace(_source(raw), sha256="0" * 64))
    assert wrong_digest.status == "failed" and wrong_digest.sheets == ()


def test_empty_and_mislabeled_bytes_do_not_yield_table() -> None:
    book = Workbook()
    output = BytesIO()
    book.save(output)
    book.close()
    raw = output.getvalue()
    empty = project_xlsx_tables(raw, _source(raw))
    assert empty.status == "deferred" and empty.reason == "empty_document"
    invalid = b"not an XLSX"
    failed = project_xlsx_tables(invalid, _source(invalid))
    assert failed.status == "failed" and failed.reason == "invalid_document"


def test_separator_budget_matches_existing_indexed_text_refusal() -> None:
    book = Workbook()
    sheet = book.active
    for row_number in range(150):
        sheet.append([
            hashlib.sha256(f"{row_number}:{column}".encode()).hexdigest()[:27]
            for column in range(50)
        ])
    output = BytesIO()
    book.save(output)
    book.close()
    raw = output.getvalue()
    source = _source(raw)
    indexed = extract_document(raw, source)
    projection = project_xlsx_tables(raw, source)
    assert indexed.status == "deferred" and indexed.reason == "text_limit"
    assert projection.status == "deferred" and projection.reason == "text_limit"
    assert projection.sheets == ()


def test_closed_pipe_contract_rejects_forged_source_rows_and_status() -> None:
    raw = _workbook()
    source = _source(raw)
    response = process_request(encode_request(
        raw, source.version_id, source.locator, source.sha256, SHEET_TYPE,
    ))
    result = decode_projection(response, source)
    assert result == project_xlsx_tables(raw, source)
    assert encode_projection(result) == response
    payload = json.loads(response)
    for mutation in (
        {"source": {**payload["source"], "sha256": "0" * 64}},
        {"scope": "complete_workbook"},
        {"status": "complete"},
        {"cell_count": 0},
        {"sheets": [{**payload["sheets"][0], "index": 2}]},
        {"sheets": [{**payload["sheets"][0], "rows": [
            {**payload["sheets"][0]["rows"][0], "number": True},
        ]}]},
        {"sheets": [{**payload["sheets"][0], "rows": [
            {**payload["sheets"][0]["rows"][0], "values": ["bad\x00cell"]},
        ]}]},
    ):
        with pytest.raises(DocumentUnavailable):
            decode_projection(json.dumps({**payload, **mutation}).encode(), source)


def test_pinned_isolated_caller_uses_bounded_pipes_and_exact_child(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw = _workbook()
    source = _source(raw)
    request = encode_request(raw, source.version_id, source.locator, source.sha256, SHEET_TYPE)
    response = process_request(request)
    processor, calls, _ = fake_docker(
        monkeypatch, processor=DockerXlsxTableProcessor(IMAGE), output=response,
    )
    result = processor.extract(
        raw, version_id=source.version_id, locator=source.locator,
        sha256=source.sha256, media_type=SHEET_TYPE,
    )
    assert result.status == "extracted"
    assert processor.processor_version == "sandbox-xlsx-tables-v1-" + "a" * 64
    create = next(args for args, _, _ in calls if args[:2] == ["container", "create"])
    assert create[-5:] == [IMAGE, "-I", "-B", "-m", "masterbrain.datacore.xlsx_tables_worker"]
    assert all("Fictional-0" not in arg for arg in create)
    assert [(args, data) for args, data, _ in calls if data] == [
        (["container", "start", "--attach", "--interactive", CID], request),
    ]
    assert calls[-1][0] == ["container", "rm", "--force", CID]


def test_substituted_table_child_result_fails_closed(monkeypatch: pytest.MonkeyPatch) -> None:
    raw = _workbook()
    source = _source(raw)
    request = encode_request(raw, source.version_id, source.locator, source.sha256, SHEET_TYPE)
    response = json.loads(process_request(request))
    response["source"]["sha256"] = "0" * 64
    processor, calls, _ = fake_docker(
        monkeypatch, processor=DockerXlsxTableProcessor(IMAGE),
        output=json.dumps(response).encode(),
    )
    with pytest.raises(DocumentUnavailable):
        processor.extract(
            raw, version_id=source.version_id, locator=source.locator,
            sha256=source.sha256, media_type=SHEET_TYPE,
        )
    assert calls[-1][0] == ["container", "rm", "--force", CID]
