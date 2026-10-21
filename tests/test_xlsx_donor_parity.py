"""Fictional XLSX goldens and optional exact-pinned FOS-Master replay.

This is a nonpublishing comparison, not workbook-feature or financial parity.
"""

from __future__ import annotations

import ast
import hashlib
import subprocess
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from masterbrain.datacore.document_contract import SHEET_TYPE, DocumentUnavailable, source_for
from masterbrain.datacore.document_worker import extract_document
from masterbrain.datacore.xlsx_tables import project_xlsx_tables

DONOR_COMMIT = "218d5caec2f20b9047fd3eb8e027dc671acf0c15"
DONOR_SHA256 = "746932e5ce033a9b8f8bd7155a88dfbb6d1c57a9916aea2230e6afbe1a3ab269"
DONOR_PATH = (
    Path(__file__).resolve().parents[2] / "FOS-Master" / "_SYSTEM" / "tools" / "doc_extract.py"
)
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _workbook(*, formula: bool = False, row_count: int = 2) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Rent roll"
    sheet.append(["Project", "Monthly USD"])
    for number in range(row_count):
        sheet.append([f"Fictional-{number}", number + 100])
    if formula:
        sheet.cell(row=2, column=2).value = "=100+1"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _datacore(raw: bytes):
    digest = hashlib.sha256(raw).hexdigest()
    source = source_for(raw, "version_" + digest[:32], "rent.xlsx", digest, SHEET_TYPE)
    return extract_document(raw, source)


def _pinned_donor_extract():
    if not DONOR_PATH.is_file():
        pytest.skip("optional FOS-Master donor checkout unavailable")
    checkout = DONOR_PATH.parents[2]
    commit = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=checkout, check=True,
        capture_output=True, text=True,
    ).stdout.strip()
    assert commit == DONOR_COMMIT, "FOS-Master donor checkout is not at the pinned commit"
    source = DONOR_PATH.read_bytes()
    assert hashlib.sha256(source).hexdigest() == DONOR_SHA256, (
        f"FOS-Master donor changed; re-pin from commit {DONOR_COMMIT} before replay"
    )
    tree = ast.parse(source, filename=str(DONOR_PATH))
    functions = [
        item for item in tree.body
        if isinstance(item, ast.FunctionDef) and item.name == "extract_xlsx"
    ]
    assert len(functions) == 1
    namespace = {
        "Path": Path, "datetime": datetime, "hashlib": hashlib, "OLE2_MAGIC": OLE2_MAGIC,
    }
    exec(compile(ast.Module(body=functions, type_ignores=[]), str(DONOR_PATH), "exec"), namespace)
    return namespace["extract_xlsx"]


def test_datacore_plain_workbook_golden_has_all_authored_values() -> None:
    result = _datacore(_workbook(row_count=3))
    assert result.status == "extracted" and result.reason == "bounded_xlsx_v1"
    assert result.metadata.sheet_count == 1
    assert result.text == (
        "[Sheet: Rent roll]\nProject\tMonthly USD\n"
        "Fictional-0\t100\nFictional-1\t101\nFictional-2\t102"
    )


def test_pinned_donor_plain_table_values_are_represented_but_shape_differs(
    tmp_path: Path,
) -> None:
    donor_extract = _pinned_donor_extract()
    raw = _workbook(row_count=3)
    path = tmp_path / "rent.xlsx"
    path.write_bytes(raw)
    donor = donor_extract(path, {"xlsx_max_rows": 2000})
    shared = _datacore(raw)
    structured = project_xlsx_tables(raw, shared.source)
    assert donor["tables"] == [{
        "sheet": "Rent roll", "headers": ["Project", "Monthly USD"],
        "rows": [["Fictional-0", "100"], ["Fictional-1", "101"], ["Fictional-2", "102"]],
    }]
    assert donor["sheets"] == [{"name": "Rent roll", "row_count": 4}]
    rows = structured.sheets[0].rows
    assert structured.status == "extracted"
    assert donor["tables"][0]["headers"] == list(rows[0].values)
    assert donor["tables"][0]["rows"] == [list(row.values) for row in rows[1:]]
    assert all(cell in shared.text for row in [
        donor["tables"][0]["headers"], *donor["tables"][0]["rows"]
    ] for cell in row)
    assert donor["full_text"] != shared.text  # donor uses spaces and tables; DataCore uses tabs.


def test_formula_cache_gap_is_explicit_and_datacore_refuses_completion(
    tmp_path: Path,
) -> None:
    donor_extract = _pinned_donor_extract()
    raw = _workbook(formula=True)
    path = tmp_path / "formula.xlsx"
    path.write_bytes(raw)
    donor = donor_extract(path, {"xlsx_max_rows": 2000})
    shared = _datacore(raw)
    assert donor["tables"][0]["rows"][0] == ["Fictional-0", ""]
    assert shared.status == "deferred" and shared.reason == "unsupported_features"
    assert shared.text == ""


def test_donor_truncation_and_datacore_full_or_refusal_are_distinct(
    tmp_path: Path,
) -> None:
    donor_extract = _pinned_donor_extract()
    raw = _workbook(row_count=50)
    path = tmp_path / "many.xlsx"
    path.write_bytes(raw)
    donor = donor_extract(path, {"xlsx_max_rows": 10})
    shared = _datacore(raw)
    assert donor["metadata"]["truncated"] is True
    assert donor["metadata"]["total_rows"] == 51
    assert donor["metadata"]["extracted_rows"] == 10
    assert len(donor["tables"][0]["rows"]) == 9
    assert "Fictional-49" not in donor["full_text"]
    assert shared.status == "extracted" and "Fictional-49" in shared.text


def test_thousand_row_plain_workbook_replays_without_losing_tail(tmp_path: Path) -> None:
    donor_extract = _pinned_donor_extract()
    raw = _workbook(row_count=1000)
    path = tmp_path / "stress.xlsx"
    path.write_bytes(raw)
    donor = donor_extract(path, {"xlsx_max_rows": 2000})
    shared = _datacore(raw)
    structured = project_xlsx_tables(raw, shared.source)
    assert len(donor["tables"][0]["rows"]) == 1000
    assert donor["tables"][0]["rows"][-1] == ["Fictional-999", "1099"]
    assert donor["tables"][0]["rows"] == [
        list(row.values) for row in structured.sheets[0].rows[1:]
    ]
    assert shared.status == "extracted"
    assert shared.metadata.cell_count == 2002
    assert shared.text.endswith("Fictional-999\t1099")


@pytest.mark.parametrize("case", range(24))
def test_pinned_donor_sparse_multisheet_value_matrix(tmp_path: Path, case: int) -> None:
    """Compare authored values, not FOS's header inference or text formatting."""
    donor_extract = _pinned_donor_extract()
    book = Workbook()
    first = book.active
    first.title = "Fictional A"
    sheets = [first]
    if case % 3 != 0:
        second = book.create_sheet("Fictional B")
        if case % 2 == 0:
            second.sheet_state = "hidden"
        sheets.append(second)
    for sheet_index, sheet in enumerate(sheets):
        if case % 4 == 0:
            sheet.append([None, None, None])
        sheet.append(["Name", "Amount", "Note"])
        for number in range(1 + case % 5):
            if (number + case) % 3 == 0:
                sheet.append([None, None, None])
            sheet.append([
                f"Fictional-{case}-{sheet_index}-{number}",
                100 + number,
                None if (case + number) % 2 else "reviewed",
            ])
    output = BytesIO()
    book.save(output)
    book.close()
    raw = output.getvalue()
    path = tmp_path / f"matrix-{case}.xlsx"
    path.write_bytes(raw)
    donor = donor_extract(path, {"xlsx_max_rows": 2000})
    source = _datacore(raw)
    projection = project_xlsx_tables(raw, source.source)
    assert source.status == "extracted" and projection.status == "extracted"
    assert len(donor["tables"]) == len(projection.sheets)
    assert donor["metadata"].get("truncated") is not True
    for donor_table, shared_sheet in zip(donor["tables"], projection.sheets, strict=True):
        assert donor_table["sheet"] == shared_sheet.name
        assert donor_table["headers"] == list(shared_sheet.rows[0].values)
        assert donor_table["rows"] == [
            list(row.values) for row in shared_sheet.rows[1:]
        ]
        assert all(row.number >= 1 for row in shared_sheet.rows)
        assert len({row.number for row in shared_sheet.rows}) == len(shared_sheet.rows)


def test_exact_row_cap_keeps_tail_and_next_row_refuses_partial_claim(tmp_path: Path) -> None:
    donor_extract = _pinned_donor_extract()
    at_cap = _workbook(row_count=1999)  # header plus 1,999 rows = 2,000 physical rows
    at_path = tmp_path / "at-cap.xlsx"
    at_path.write_bytes(at_cap)
    donor_at = donor_extract(at_path, {"xlsx_max_rows": 2000})
    shared_at = _datacore(at_cap)
    rows_at = project_xlsx_tables(at_cap, shared_at.source)
    assert donor_at["metadata"].get("truncated") is not True
    assert shared_at.status == "extracted" and rows_at.status == "extracted"
    assert rows_at.sheets[0].rows[-1].number == 2000
    assert donor_at["tables"][0]["rows"][-1] == list(rows_at.sheets[0].rows[-1].values)

    over_cap = _workbook(row_count=2000)
    over_path = tmp_path / "over-cap.xlsx"
    over_path.write_bytes(over_cap)
    donor_over = donor_extract(over_path, {"xlsx_max_rows": 2000})
    shared_over = _datacore(over_cap)
    rows_over = project_xlsx_tables(over_cap, shared_over.source)
    assert donor_over["metadata"]["truncated"] is True
    assert donor_over["metadata"]["total_rows"] == 2001
    assert donor_over["metadata"]["extracted_rows"] == 2000
    assert shared_over.status == "deferred" and shared_over.reason == "structure_limit"
    assert shared_over.text == ""
    assert rows_over.status == "deferred" and rows_over.reason == "structure_limit"
    assert rows_over.sheets == () and rows_over.cell_count == 0


def test_encrypted_or_mislabeled_workbook_never_produces_text(tmp_path: Path) -> None:
    donor_extract = _pinned_donor_extract()
    for name, raw in (("encrypted.xlsx", OLE2_MAGIC + b"\x00" * 100),
                      ("misnamed.xlsx", b"not a workbook")):
        path = tmp_path / name
        path.write_bytes(raw)
        donor = donor_extract(path, {"xlsx_max_rows": 2000})
        shared = _datacore(raw)
        assert "error" in donor and not donor["full_text"]
        assert shared.status != "extracted" and shared.text == ""
        if name == "encrypted.xlsx":
            assert donor["password_protected"] is True
            assert shared.reason == "invalid_document"  # Different error vocabulary.


def test_macro_enabled_donor_extension_is_not_silently_accepted_by_datacore(
    tmp_path: Path,
) -> None:
    donor_extract = _pinned_donor_extract()
    raw = _workbook()
    path = tmp_path / "macro.xlsm"
    path.write_bytes(raw)
    donor = donor_extract(path, {"xlsx_max_rows": 2000})
    assert donor["type"] == "xlsx" and donor["tables"]
    digest = hashlib.sha256(raw).hexdigest()
    with pytest.raises(DocumentUnavailable):
        source_for(raw, "version_" + digest[:32], "macro.xlsm", digest, SHEET_TYPE)
