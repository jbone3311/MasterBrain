"""Independently read actual synthetic formats and verify the authored oracle."""

import hashlib
import json
import os
import socket
import struct
import subprocess
import sys
import zlib
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from email_brain.store import prepare_mime
from email_brain.synthetic_formats import SHEET_TYPE, WORD_TYPE
from email_brain.synthetic_mailbox import (
    MailboxSpec,
    export_mailbox,
    mailbox_manifest,
    project_events,
    project_facts,
)


def test_versioned_seed_42_manifest_golden():
    assert mailbox_manifest(MailboxSpec(42, 1))["manifest_sha256"] == (
        "4c6e8d82850919607f4bf4623032ff147c319cd743a0ba840445afad766f9f3f"
    )


@pytest.mark.parametrize("seed", [0, 1, 42, 9123, 2**31 - 1])
def test_reproducible_world_manifest_and_current_parser_oracle(seed):
    spec = MailboxSpec(seed, 2)
    manifest = mailbox_manifest(spec)
    assert manifest == mailbox_manifest(spec)
    digest = manifest.pop("manifest_sha256")
    canonical = json.dumps(
        manifest, sort_keys=True, separators=(",", ":"), ensure_ascii=True
    ).encode()
    assert hashlib.sha256(canonical).hexdigest() == digest
    assert manifest["expected_final_counts"] == {
        "scheduled_events": 30,
        "unique_provider_events": 28,
        "ready": 20,
        "quarantined": 2,
        "deleted": 2,
        "source_files": 24,
    }
    for project in range(2):
        events = project_events(seed, project)
        assert events == project_events(seed, project)
        assert len(events) == 15 and len({e.item.event_key for e in events}) == 14
        assert len({e.filename for e in events if e.filename}) == 12
        assert list(e.arrival_at for e in events) == sorted(e.arrival_at for e in events)
        assert events[1].item.received_at > events[2].item.received_at
        assert events[12].item == events[2].item
        assert events[13].item.raw == events[0].item.raw
        assert events[13].item.keywords == ("$seen",)
        assert events[14].item.raw is None and events[14].kind == "tombstone"
        for event in events:
            if event.kind != "created":
                continue
            normalized, attachments, quarantined = prepare_mime(event.item.raw)
            assert quarantined == (event.expected_state == "quarantined")
            if quarantined:
                assert normalized is None and not attachments
                continue
            assert event.marker in normalized.text
            assert normalized.raw_sha256 == hashlib.sha256(event.item.raw).hexdigest()
            assert normalized.sender.endswith("@example.invalid")
            assert len(attachments) == len(event.attachments)
            for actual, expected in zip(attachments, event.attachments, strict=True):
                assert (actual.original_filename, actual.media_type) == (
                    expected.name,
                    expected.media_type,
                )
                assert (actual.status, actual.reason) == (
                    expected.expected_status,
                    expected.expected_reason,
                )
                if expected.media_type != "message/rfc822":
                    assert actual.sha256 == hashlib.sha256(expected.raw).hexdigest()
                if expected.expected_status == "extracted":
                    assert expected.known_text in actual.text
                else:
                    assert actual.text == ""


def test_seed_changes_fact_world_and_bytes_without_random_or_network(monkeypatch):
    monkeypatch.setattr(socket, "create_connection", Mock(side_effect=AssertionError("network")))
    monkeypatch.setattr(socket.socket, "connect", Mock(side_effect=AssertionError("network")))
    assert mailbox_manifest(MailboxSpec(42, 1)) != mailbox_manifest(MailboxSpec(43, 1))
    assert len({project_facts(seed, 0)["rent_v2_usd"] for seed in range(20)}) > 15
    assert all(
        project_facts(seed, 0)["semantic_targets"]["evaluated"] is False for seed in range(20)
    )


@pytest.mark.parametrize(
    "seed,projects",
    [(-1, 1), (2**31, 1), (True, 1), ("1", 1), (1, 0), (1, 501), (1, True), (1, 1.5)],
)
def test_generator_bounds(seed, projects):
    with pytest.raises(ValueError):
        MailboxSpec(seed, projects)


def test_independent_pdf_office_and_png_readers():
    facts = project_facts(42, 0)
    for event in project_events(42, 0):
        if event.kind != "created":
            continue
        for sample in event.attachments:
            if sample.media_type == "application/pdf":
                if not sample.well_formed:
                    with pytest.raises(PdfReadError):
                        PdfReader(BytesIO(sample.raw), strict=True)
                    continue
                reader = PdfReader(BytesIO(sample.raw), strict=True)
                assert len(reader.pages) == 1 and not reader.is_encrypted
                assert "/OpenAction" not in reader.trailer["/Root"]
                assert "/Names" not in reader.trailer["/Root"]
                text = reader.pages[0].extract_text()
                if sample.name == "image-only-plan.pdf":
                    assert text == ""
                    image = reader.pages[0]["/Resources"]["/XObject"]["/Im1"]
                    assert image["/Width"] == 64 and image["/Height"] == 32
                    assert len(image.get_data()) == 6144
                elif sample.name == "overview.pdf":
                    assert str(facts["rent_v1_usd"]) in text and sample.known_text in text
                else:
                    cents = facts["invoice_cents"]
                    assert f"{cents // 100}.{cents % 100:02d}" in text
            elif sample.media_type == SHEET_TYPE:
                workbook = load_workbook(BytesIO(sample.raw), read_only=True, keep_links=False)
                try:
                    rows = list(workbook.active.values)
                    revision = 1 if "v1" in sample.name else 2
                    assert rows[1][:3] == (
                        facts["project"],
                        facts[f"rent_v{revision}_usd"],
                        revision,
                    )
                finally:
                    workbook.close()
            elif sample.media_type == WORD_TYPE:
                text = "\n".join(p.text for p in Document(BytesIO(sample.raw)).paragraphs)
                assert text == sample.known_text and str(facts["rent_v2_usd"]) in text
            elif sample.media_type == "image/png":
                assert sample.raw[:8] == b"\x89PNG\r\n\x1a\n"
                offset, chunks = 8, {}
                while offset < len(sample.raw):
                    size = struct.unpack(">I", sample.raw[offset : offset + 4])[0]
                    kind = sample.raw[offset + 4 : offset + 8]
                    data = sample.raw[offset + 8 : offset + 8 + size]
                    crc = struct.unpack(">I", sample.raw[offset + 8 + size : offset + 12 + size])[0]
                    assert zlib.crc32(kind + data) == crc
                    chunks[kind] = data
                    offset += 12 + size
                assert set(chunks) == {b"IHDR", b"IDAT", b"IEND"}
                assert struct.unpack(">II", chunks[b"IHDR"][:8]) == (64, 32)
                decoded = zlib.decompress(chunks[b"IDAT"])
                assert len(decoded) == 32 * (1 + 64 * 3)
                assert all(decoded[row * 193] == 0 for row in range(32))
            if sample.media_type in (WORD_TYPE, SHEET_TYPE, "application/zip"):
                with ZipFile(BytesIO(sample.raw)) as archive:
                    assert archive.testzip() is None
                    for member in archive.infolist():
                        assert (
                            not member.filename.startswith("/")
                            and ".." not in Path(member.filename).parts
                        )
                        assert member.date_time == (2026, 1, 1, 0, 0, 0)
                        assert b'TargetMode="External"' not in archive.read(member)
                        assert "vba" not in member.filename.lower()


def test_reply_headers_nested_source_and_calendar_revisions():
    facts = project_facts(42, 0)
    events = project_events(42, 0)
    ids = {event.message_id for event in events if event.message_id}
    calendars = []
    for event in events:
        if event.kind != "created" or not event.message_id:
            continue
        mail = BytesParser(policy=policy.SMTP).parsebytes(event.item.raw)
        assert mail["Message-ID"] == event.message_id
        if event.parent_id:
            assert event.parent_id in ids
            assert mail["In-Reply-To"] == event.parent_id and event.parent_id in mail["References"]
        for sample in event.attachments:
            if sample.media_type == "text/calendar":
                source = sample.raw.decode()
                assert f"UID:{facts['meeting_uid']}\r\n" in source
                assert all(len(line) <= 75 for line in sample.raw.split(b"\r\n"))
                calendars.append(source)
            elif sample.media_type == "message/rfc822":
                nested = next(
                    part for part in mail.walk() if part.get_content_type() == "message/rfc822"
                )
                assert nested.get_payload()[0].as_bytes(policy=policy.SMTP) == events[0].item.raw
    assert "SEQUENCE:0" in calendars[0] and "SEQUENCE:1" in calendars[1]
    assert "DTSTART:20260905T090000Z" in calendars[0]
    assert "DTSTART:20260906T090000Z" in calendars[1]


def test_packet_export_is_complete_exclusive_and_matches_manifest(tmp_path):
    path = tmp_path.resolve() / "packet"
    result = export_mailbox(path, MailboxSpec(42, 2))
    assert result["source_files"] == 24
    manifest = json.loads((path / "manifest.json").read_text())
    assert manifest["manifest_sha256"] == result["manifest_sha256"]
    assert len(list(path.iterdir())) == 25
    for row in manifest["events"]:
        if row["file"]:
            file = path / row["file"]
            assert hashlib.sha256(file.read_bytes()).hexdigest() == row["raw_sha256"]
            assert file.stat().st_mode & 0o777 == 0o600
    assert path.stat().st_mode & 0o777 == 0o700
    before = {p.name: p.read_bytes() for p in path.iterdir()}
    with pytest.raises(FileExistsError):
        export_mailbox(path, MailboxSpec(43, 1))
    assert before == {p.name: p.read_bytes() for p in path.iterdir()}


@pytest.mark.parametrize(
    "kind", ["relative", "traversal", "dropbox", "symlink", "existing_file", "source"]
)
def test_export_refuses_unsafe_or_existing_targets(tmp_path, kind):
    base = tmp_path.resolve()
    path = base / "new"
    if kind == "relative":
        path = Path("relative")
    elif kind == "traversal":
        path = base / ".." / "new"
    elif kind == "dropbox":
        path = base / "Dropbox" / "new"
    elif kind == "symlink":
        (base / "alias").symlink_to(base, target_is_directory=True)
        path = base / "alias" / "new"
    elif kind == "existing_file":
        path.write_text("preserve")
    else:
        path = Path(__file__).resolve().parents[1] / "generated-never-created"
    with pytest.raises((ValueError, OSError)):
        export_mailbox(path, MailboxSpec(42, 1))
    if kind == "existing_file":
        assert path.read_text() == "preserve"
    else:
        assert not (base / "new").exists()


def test_failed_packet_has_no_completion_manifest(tmp_path, monkeypatch):
    real = os.fsync
    calls = 0

    def failure(fd):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("simulated disk failure")
        return real(fd)

    monkeypatch.setattr(os, "fsync", failure)
    path = tmp_path.resolve() / "partial"
    with pytest.raises(OSError):
        export_mailbox(path, MailboxSpec(42, 1))
    assert path.is_dir() and not (path / "manifest.json").exists()


def test_cli_export_needs_no_database_and_has_content_free_summary(tmp_path):
    path = tmp_path.resolve() / "cli-packet"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "email_brain.cli",
            "export-test-mailbox",
            "--output-dir",
            str(path),
            "--projects",
            "1",
        ],
        env={"PATH": "/usr/bin:/bin"},
        text=True,
        capture_output=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0 and not result.stderr
    assert json.loads(result.stdout)["source_files"] == 12
    assert "broker" not in result.stdout and str(path) not in result.stdout
