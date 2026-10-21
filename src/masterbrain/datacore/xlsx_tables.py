"""Worker-only, nonpublishing XLSX row and cell projection.

This additive path preserves the existing `bounded_xlsx_v1` text result. It
does not identify financial facts or certify the whole workbook's features.
Only the isolated document worker may run this on untrusted source bytes.
"""

from __future__ import annotations

import hashlib
import hmac
import json
import re
from dataclasses import asdict, dataclass
from io import BytesIO

from masterbrain.datacore.document_contract import (
    MAX_CELLS,
    MAX_COLUMNS,
    MAX_DOCUMENT_BYTES,
    MAX_ROWS_PER_SHEET,
    MAX_SHEETS,
    MAX_TEXT_CHARS,
    DocumentSource,
    DocumentUnavailable,
    _fields,
    _load,
)
from masterbrain.datacore.document_worker import (
    _cell_text,
    _CellTextLimit,
    _zip_preflight,
)


@dataclass(frozen=True, slots=True)
class XlsxRow:
    number: int
    values: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class XlsxSheet:
    index: int
    name: str
    visibility: str
    rows: tuple[XlsxRow, ...]


@dataclass(frozen=True, slots=True)
class XlsxTableProjection:
    source: DocumentSource
    status: str
    reason: str
    sheets: tuple[XlsxSheet, ...] = ()
    cell_count: int = 0
    scope: str = "bounded_textual_cell_values_only"


CONTRACT = "sandbox-xlsx-tables-v1"
MAX_OUTPUT_BYTES = 1024 * 1024
_OUTCOMES = {
    "bounded_xlsx_tables_v1": "extracted",
    "empty_document": "deferred",
    "encrypted": "deferred",
    "unsupported_features": "deferred",
    "structure_limit": "deferred",
    "text_limit": "deferred",
    "invalid_document": "failed",
}
_SHA256 = re.compile(r"^[a-f0-9]{64}$")


def _indexed_text(sheets: tuple[XlsxSheet, ...]) -> str:
    return "\n\n".join(
        f"[Sheet: {sheet.name}]\n"
        + "\n".join("\t".join(row.values) for row in sheet.rows)
        for sheet in sheets
    ).strip()


def require_indexed_text_digest(
    projection: XlsxTableProjection, expected_sha256: str,
) -> str:
    """Bind cells to an already-known DataCore document body digest.

    This checks the exact existing `bounded_xlsx_v1` text layout. It does not
    independently prove a child parsed the source bytes or authorize a read.
    """
    if (
        type(projection) is not XlsxTableProjection
        or projection.status != "extracted"
        or projection.reason != "bounded_xlsx_tables_v1"
        or type(expected_sha256) is not str
        or _SHA256.fullmatch(expected_sha256) is None
    ):
        raise DocumentUnavailable()
    text = _indexed_text(projection.sheets)
    if not text or len(text) > MAX_TEXT_CHARS:
        raise DocumentUnavailable()
    digest = hashlib.sha256(text.encode("utf-8", errors="strict")).hexdigest()
    if not hmac.compare_digest(digest, expected_sha256):
        raise DocumentUnavailable()
    return digest


def projection_layout_sha256(projection: XlsxTableProjection) -> str:
    """Digest the extracted sheet/row layout independently of source packaging."""
    if (
        type(projection) is not XlsxTableProjection
        or projection.status != "extracted"
        or projection.reason != "bounded_xlsx_tables_v1"
        or not projection.sheets
    ):
        raise DocumentUnavailable()
    payload = json.dumps(
        [asdict(sheet) for sheet in projection.sheets],
        ensure_ascii=True, allow_nan=False, separators=(",", ":"),
    ).encode("ascii")
    if len(payload) > MAX_OUTPUT_BYTES:
        raise DocumentUnavailable()
    return hashlib.sha256(payload).hexdigest()


def encode_projection(result: XlsxTableProjection) -> bytes:
    """Serialize the closed worker result without changing document-v1 output."""
    if type(result) is not XlsxTableProjection:
        raise DocumentUnavailable()
    return json.dumps(
        {"contract_version": CONTRACT, **asdict(result)},
        ensure_ascii=True, allow_nan=False, separators=(",", ":"),
    ).encode("ascii")


def decode_projection(payload: bytes, source: DocumentSource) -> XlsxTableProjection:
    """Treat every byte returned by a child as untrusted, including its source."""
    try:
        value = _fields(
            _load(payload, MAX_OUTPUT_BYTES),
            {"contract_version", "source", "status", "reason", "sheets", "cell_count", "scope"},
        )
        if (
            type(source) is not DocumentSource
            or source.format != "xlsx"
            or value["contract_version"] != CONTRACT
            or _fields(value["source"], set(asdict(source))) != asdict(source)
            or value["scope"] != "bounded_textual_cell_values_only"
            or type(value["reason"]) is not str
            or value["reason"] not in _OUTCOMES
            or value["status"] != _OUTCOMES[value["reason"]]
            or type(value["cell_count"]) is not int
            or not 0 <= value["cell_count"] <= MAX_CELLS
            or type(value["sheets"]) is not list
            or len(value["sheets"]) > MAX_SHEETS
        ):
            raise DocumentUnavailable()
        sheets: list[XlsxSheet] = []
        represented_cells = text_chars = nonempty = 0
        for index, sheet_data in enumerate(value["sheets"], 1):
            sheet = _fields(sheet_data, {"index", "name", "visibility", "rows"})
            name, rows_data = sheet["name"], sheet["rows"]
            if (
                type(sheet["index"]) is not int
                or sheet["index"] != index
                or type(name) is not str
                or not 1 <= len(name) <= 200
                or "\x00" in name
                or sheet["visibility"] not in {"visible", "hidden", "veryHidden"}
                or type(rows_data) is not list
                or not 1 <= len(rows_data) <= MAX_ROWS_PER_SHEET
            ):
                raise DocumentUnavailable()
            name.encode("utf-8", errors="strict")
            rows: list[XlsxRow] = []
            previous = 0
            for row_data in rows_data:
                row = _fields(row_data, {"number", "values"})
                number, values = row["number"], row["values"]
                if (
                    type(number) is not int
                    or not previous < number <= MAX_ROWS_PER_SHEET
                    or type(values) is not list
                    or not 1 <= len(values) <= MAX_COLUMNS
                    or not all(type(item) is str for item in values)
                ):
                    raise DocumentUnavailable()
                previous = number
                for item in values:
                    if len(item) > 1000 or "\x00" in item:
                        raise DocumentUnavailable()
                    item.encode("utf-8", errors="strict")
                    text_chars += len(item)
                    nonempty += bool(item)
                represented_cells += len(values)
                if not any(values):
                    raise DocumentUnavailable()
                rows.append(XlsxRow(number, tuple(values)))
            sheets.append(XlsxSheet(index, name, sheet["visibility"], tuple(rows)))
        if (
            text_chars > MAX_TEXT_CHARS
            or represented_cells > value["cell_count"]
            or (
                value["status"] == "extracted"
                and (not sheets or nonempty == 0 or value["cell_count"] == 0)
            )
            or (value["status"] != "extracted" and (sheets or value["cell_count"] != 0))
        ):
            raise DocumentUnavailable()
        return XlsxTableProjection(
            source, value["status"], value["reason"], tuple(sheets),
            value["cell_count"], value["scope"],
        )
    except Exception:
        raise DocumentUnavailable() from None


def _outcome(source: DocumentSource, status: str, reason: str) -> XlsxTableProjection:
    return XlsxTableProjection(source, status, reason)


def project_xlsx_tables(raw: bytes, source: DocumentSource) -> XlsxTableProjection:
    """Capture physical rows, positions and values; never infer a true header.

    Every nonempty value is addressable by `(sheet.index, row.number, column
    index)`. Empty cells stay in each row's positional tuple. No partial table
    is returned after formula, count, text or OOXML safety refusal.
    """
    if (
        type(raw) is not bytes
        or type(source) is not DocumentSource
        or source.format != "xlsx"
        or len(raw) != source.size_bytes
        or len(raw) > MAX_DOCUMENT_BYTES
        or not hmac.compare_digest(hashlib.sha256(raw).hexdigest(), source.sha256)
    ):
        return _outcome(source, "failed", "invalid_document")
    reason = _zip_preflight(raw, "xl/workbook.xml")
    if reason:
        return _outcome(source, "failed" if reason == "invalid_document" else "deferred", reason)
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.sheetnames) > MAX_SHEETS:
                return _outcome(source, "deferred", "structure_limit")
            sheets: list[XlsxSheet] = []
            cell_count = text_chars = nonempty = 0
            for index, sheet in enumerate(book.worksheets, 1):
                if sheet.sheet_state not in {"visible", "hidden", "veryHidden"}:
                    return _outcome(source, "failed", "invalid_document")
                max_row, max_col = sheet.max_row, sheet.max_column
                if (isinstance(max_row, int) and max_row > MAX_ROWS_PER_SHEET) or (
                    isinstance(max_col, int) and max_col > MAX_COLUMNS
                ):
                    return _outcome(source, "deferred", "structure_limit")
                # Read-only iteration is bounded by the declared <dimension>; a
                # writer that understates it would hide cells (and formulas)
                # from an extraction still reported complete. Measure the real
                # extent and recheck it; a truthful dimension reads identically.
                sheet.reset_dimensions()
                if next(sheet.iter_rows(min_row=1, min_col=1), None) is not None:
                    sheet.calculate_dimension(force=True)  # fails on an empty sheet
                max_row, max_col = sheet.max_row, sheet.max_column
                if (isinstance(max_row, int) and max_row > MAX_ROWS_PER_SHEET) or (
                    isinstance(max_col, int) and max_col > MAX_COLUMNS
                ):
                    return _outcome(source, "deferred", "structure_limit")
                rows: list[XlsxRow] = []
                for row_number, row in enumerate(
                    sheet.iter_rows(min_row=1, min_col=1, values_only=False), 1,
                ):
                    if row_number > MAX_ROWS_PER_SHEET or len(row) > MAX_COLUMNS:
                        return _outcome(source, "deferred", "structure_limit")
                    if any(cell.data_type == "f" for cell in row):
                        return _outcome(source, "deferred", "unsupported_features")
                    values = tuple(_cell_text(cell.value) for cell in row)
                    cell_count += len(values)
                    text_chars += sum(len(value) for value in values)
                    if cell_count > MAX_CELLS:
                        return _outcome(source, "deferred", "structure_limit")
                    if text_chars > MAX_TEXT_CHARS:
                        return _outcome(source, "deferred", "text_limit")
                    if any(values):
                        nonempty += 1
                        rows.append(XlsxRow(row_number, values))
                if rows:
                    sheets.append(
                        XlsxSheet(index, sheet.title[:200], sheet.sheet_state, tuple(rows))
                    )
            if nonempty == 0:
                return _outcome(source, "deferred", "empty_document")
            result = XlsxTableProjection(
                source, "extracted", "bounded_xlsx_tables_v1", tuple(sheets), cell_count,
            )
            if len(_indexed_text(result.sheets)) > MAX_TEXT_CHARS:
                return _outcome(source, "deferred", "text_limit")
            return result
        finally:
            book.close()
    except _CellTextLimit:
        return _outcome(source, "deferred", "text_limit")
    except Exception:
        return _outcome(source, "failed", "invalid_document")
