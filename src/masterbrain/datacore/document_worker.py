"""Restricted pipe-only FileOS PDF/DOCX/XLSX extractor."""

from __future__ import annotations

import logging
import stat
import sys
from io import BytesIO
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from email_brain.mime_sandbox import _resource_limits
from masterbrain.datacore.document_contract import (
    MAX_CELLS,
    MAX_COLUMNS,
    MAX_DOCUMENT_BYTES,
    MAX_OUTPUT_BYTES,
    MAX_PAGES,
    MAX_PARAGRAPHS,
    MAX_REQUEST_BYTES,
    MAX_ROWS_PER_SHEET,
    MAX_SHEETS,
    MAX_TABLES,
    MAX_TEXT_CHARS,
    MAX_ZIP_EXPANDED_BYTES,
    MAX_ZIP_MEMBERS,
    MAX_ZIP_RATIO,
    DocumentMetadata,
    DocumentResult,
    DocumentSource,
    DocumentUnavailable,
    decode_request,
    decode_result,
    encode_result,
)


class _EncryptedPdf(Exception):
    pass


class _CellTextLimit(Exception):
    pass


def _safe_property(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\x00", "").strip()
    if not text:
        return None
    text.encode("utf-8", errors="strict")
    return text[:300]


def _outcome(
    source: DocumentSource,
    status: str,
    reason: str,
    *,
    text: str = "",
    metadata: DocumentMetadata | None = None,
) -> DocumentResult:
    return DocumentResult(source, status, reason, text, metadata or DocumentMetadata())


def _zip_preflight(raw: bytes, required: str) -> str | None:
    """Return a deferred reason, or None when the bounded OOXML package is safe to parse."""
    try:
        with ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS or required not in {
                item.filename for item in members
            }:
                return "structure_limit"
            expanded = 0
            names: set[str] = set()
            for item in members:
                name = item.filename
                if (
                    not name
                    or name in names
                    or name.startswith(("/", "\\"))
                    or "\\" in name
                    or ".." in name.split("/")
                    or item.flag_bits & 1
                    or stat.S_ISLNK(item.external_attr >> 16)
                    or name.lower().endswith("vbaproject.bin")
                ):
                    return "encrypted" if item.flag_bits & 1 else "unsupported_features"
                names.add(name)
                expanded += item.file_size
                if expanded > MAX_ZIP_EXPANDED_BYTES:
                    return "structure_limit"
                if item.file_size > max(1, item.compress_size) * MAX_ZIP_RATIO:
                    return "structure_limit"
                if name.endswith((".xml", ".rels")):
                    if item.file_size > 1024 * 1024:
                        return "structure_limit"
                    xml = archive.read(item)
                    lowered = xml.lower()
                    if b"<!doctype" in lowered or b"<!entity" in lowered:
                        return "unsupported_features"
                    if name.endswith(".rels"):
                        try:
                            root = ElementTree.fromstring(xml)
                        except ElementTree.ParseError:
                            return "invalid_document"
                        if any(
                            str(value).casefold() == "external"
                            for element in root.iter()
                            for key, value in element.attrib.items()
                            if key.rsplit("}", 1)[-1].casefold() == "targetmode"
                        ):
                            return "unsupported_features"
        return None
    except (BadZipFile, OSError, RuntimeError, ValueError):
        return "invalid_document"


def _extract_pdf(raw: bytes, source: DocumentSource) -> DocumentResult:
    if not raw.startswith(b"%PDF-"):
        return _outcome(source, "failed", "invalid_document")
    import pypdf
    from pypdf.errors import PdfReadError

    if pypdf.__version__ != "6.19.0":
        raise DocumentUnavailable()

    class NoDecryptReader(pypdf.PdfReader):
        def _handle_encryption(self, password: str | bytes | None) -> None:
            raise _EncryptedPdf()

    try:
        reader = NoDecryptReader(BytesIO(raw), strict=True)
        try:
            if reader.is_encrypted:
                return _outcome(source, "deferred", "encrypted")
            if any(
                key in reader.root_object
                for key in ("/OpenAction", "/AA", "/Names", "/AcroForm", "/AF", "/Collection")
            ):
                return _outcome(source, "deferred", "unsupported_features")
            page_count = len(reader.pages)
            if page_count == 0:
                return _outcome(source, "failed", "invalid_document")
            if page_count > MAX_PAGES:
                return _outcome(source, "deferred", "page_limit")
            parts: list[str] = []
            stream_bytes = 0
            pages_without_text = False
            for number, page in enumerate(reader.pages, 1):
                if any(key in page for key in ("/AA", "/Annots", "/AF")):
                    return _outcome(source, "deferred", "unsupported_features")
                content = page.get_contents()
                size = len(content.get_data()) if content is not None else 0
                stream_bytes += size
                if size > 2 * 1024 * 1024 or stream_bytes > 8 * 1024 * 1024:
                    return _outcome(source, "deferred", "structure_limit")
                text = page.extract_text().replace("\x00", "").strip()
                text.encode("utf-8", errors="strict")
                if text:
                    parts.append(f"[Page {number}]\n{text}")
                else:
                    pages_without_text = True
            joined = "\n\n".join(parts).strip()
            metadata: Any = reader.metadata or {}
            result_metadata = DocumentMetadata(
                title=_safe_property(metadata.get("/Title")),
                author=_safe_property(metadata.get("/Author")),
                page_count=page_count,
            )
            if not joined:
                return _outcome(source, "deferred", "empty_document", metadata=result_metadata)
            if pages_without_text:
                return _outcome(source, "deferred", "pages_without_text", metadata=result_metadata)
            if len(joined) > MAX_TEXT_CHARS:
                return _outcome(source, "deferred", "text_limit", metadata=result_metadata)
            return _outcome(
                source, "extracted", "bounded_pdf_v1", text=joined, metadata=result_metadata
            )
        finally:
            reader.close()
    except _EncryptedPdf:
        return _outcome(source, "deferred", "encrypted")
    except DocumentUnavailable:
        raise
    except PdfReadError:
        return _outcome(source, "failed", "invalid_document")
    except Exception:
        return _outcome(source, "failed", "invalid_document")


def _extract_docx(raw: bytes, source: DocumentSource) -> DocumentResult:
    reason = _zip_preflight(raw, "word/document.xml")
    if reason:
        return _outcome(source, "failed" if reason == "invalid_document" else "deferred", reason)
    try:
        from docx import Document

        document = Document(BytesIO(raw))
        if len(document.paragraphs) > MAX_PARAGRAPHS or len(document.tables) > MAX_TABLES:
            return _outcome(source, "deferred", "structure_limit")
        parts = [
            paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
        ]
        cell_count = 0
        for table in document.tables:
            rows: list[str] = []
            for row in table.rows:
                cell_count += len(row.cells)
                if cell_count > MAX_CELLS:
                    return _outcome(source, "deferred", "structure_limit")
                values = [cell.text.replace("\x00", "").strip() for cell in row.cells]
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                parts.append("\n".join(rows))
        text = "\n\n".join(parts).replace("\x00", "").strip()
        text.encode("utf-8", errors="strict")
        metadata = DocumentMetadata(
            title=_safe_property(document.core_properties.title),
            author=_safe_property(document.core_properties.author),
            paragraph_count=len(document.paragraphs),
            table_count=len(document.tables),
            cell_count=cell_count,
        )
        if not text:
            return _outcome(source, "deferred", "empty_document", metadata=metadata)
        if len(text) > MAX_TEXT_CHARS:
            return _outcome(source, "deferred", "text_limit", metadata=metadata)
        return _outcome(source, "extracted", "bounded_docx_v1", text=text, metadata=metadata)
    except Exception:
        return _outcome(source, "failed", "invalid_document")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\x00", "").strip()
    text.encode("utf-8", errors="strict")
    if len(text) > 1_000:
        raise _CellTextLimit()
    return text


def _extract_xlsx(raw: bytes, source: DocumentSource) -> DocumentResult:
    reason = _zip_preflight(raw, "xl/workbook.xml")
    if reason:
        return _outcome(source, "failed" if reason == "invalid_document" else "deferred", reason)
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        # Keep formula cells visible. data_only=True silently turns an uncached
        # formula into None, which would falsely certify an incomplete sheet.
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.sheetnames) > MAX_SHEETS:
                return _outcome(source, "deferred", "structure_limit")
            parts: list[str] = []
            cell_count = 0
            for sheet in workbook.worksheets:
                max_row = sheet.max_row
                max_column = sheet.max_column
                if (isinstance(max_row, int) and max_row > MAX_ROWS_PER_SHEET) or (
                    isinstance(max_column, int) and max_column > MAX_COLUMNS
                ):
                    return _outcome(source, "deferred", "structure_limit")
                rows: list[str] = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=False), 1):
                    if row_number > MAX_ROWS_PER_SHEET or len(row) > MAX_COLUMNS:
                        return _outcome(source, "deferred", "structure_limit")
                    if any(cell.data_type == "f" for cell in row):
                        return _outcome(source, "deferred", "unsupported_features")
                    try:
                        values = [_cell_text(cell.value) for cell in row]
                    except _CellTextLimit:
                        return _outcome(source, "deferred", "text_limit")
                    cell_count += len(values)
                    if cell_count > MAX_CELLS:
                        return _outcome(source, "deferred", "structure_limit")
                    if any(values):
                        rows.append("\t".join(values))
                if rows:
                    parts.append(f"[Sheet: {sheet.title[:200]}]\n" + "\n".join(rows))
            text = "\n\n".join(parts).strip()
            metadata = DocumentMetadata(
                title=_safe_property(workbook.properties.title),
                author=_safe_property(workbook.properties.creator),
                sheet_count=len(workbook.sheetnames),
                cell_count=cell_count,
            )
            if not text:
                return _outcome(source, "deferred", "empty_document", metadata=metadata)
            if len(text) > MAX_TEXT_CHARS:
                return _outcome(source, "deferred", "text_limit", metadata=metadata)
            return _outcome(source, "extracted", "bounded_xlsx_v1", text=text, metadata=metadata)
        finally:
            workbook.close()
    except Exception:
        return _outcome(source, "failed", "invalid_document")


def extract_document(raw: bytes, source: DocumentSource) -> DocumentResult:
    if len(raw) > MAX_DOCUMENT_BYTES:
        return _outcome(source, "deferred", "document_byte_limit")
    if source.format == "pdf":
        return _extract_pdf(raw, source)
    if source.format == "docx":
        return _extract_docx(raw, source)
    if source.format == "xlsx":
        return _extract_xlsx(raw, source)
    raise DocumentUnavailable()


def process_request(payload: bytes) -> bytes:
    try:
        raw, source = decode_request(payload)
        result = extract_document(raw, source)
        response = encode_result(result)
        if len(response) > MAX_OUTPUT_BYTES:
            response = encode_result(_outcome(source, "deferred", "text_limit"))
        decode_result(response, source)
        return response
    except Exception:
        raise DocumentUnavailable() from None


def main() -> int:
    try:
        _resource_limits()
        logging.disable(sys.maxsize)
        payload = sys.stdin.buffer.read(MAX_REQUEST_BYTES + 1)
        response = process_request(payload)
        if len(response) > MAX_OUTPUT_BYTES:
            raise DocumentUnavailable()
        sys.stdout.buffer.write(response)
        sys.stdout.buffer.flush()
        return 0
    except Exception:
        return 70


if __name__ == "__main__":
    raise SystemExit(main())
